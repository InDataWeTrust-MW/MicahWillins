#!/usr/bin/env python3
"""
Generate ATS-optimized DOCX resume for Micah Willins.
Parses Resume.html directly so DOCX content stays verbatim-accurate to the
approved source (no invented metrics, no manual duplication/drift).
"""

from __future__ import annotations

import re
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DEV_ROOT = Path(__file__).resolve().parents[1]
SOURCE = DEV_ROOT / "Resume.html"
OUTPUT = DEV_ROOT / "Micah-Willins-ATS.docx"

BRAND_BLUE = RGBColor(31, 75, 153)
LINKEDIN_URL = "https://www.linkedin.com/in/indatawetrust-mw"
RESUME_URL = "https://indatawetrust-mw.github.io/MicahWillins/"

def clean(text: str) -> str:
    text = text.replace("\u2011", "-")
    return re.sub(r"[\s\xa0]+", " ", text).strip()


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attribute), "Aptos Narrow")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "15")
    r_pr.append(size)
    r_pr.append(OxmlElement("w:i"))
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = BRAND_BLUE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True


def li_text(li: Tag) -> tuple[str, str]:
    """Return (bold_lead, rest) for a docx-outline <li>, preserving verbatim wording."""
    strong = li.find(["strong", "b"], recursive=False)
    if strong:
        lead = clean(strong.get_text())
        rest_parts = []
        started = False
        for child in li.children:
            if child is strong:
                started = True
                continue
            if not started:
                continue
            if isinstance(child, NavigableString):
                rest_parts.append(str(child))
            elif isinstance(child, Tag):
                rest_parts.append(child.get_text())
        rest = clean("".join(rest_parts))
        return lead, rest
    return "", clean(li.get_text())


def build() -> None:
    soup = BeautifulSoup(SOURCE.read_text(encoding="utf-8"), "html.parser")
    container = soup.find("div", class_="container")

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.core_properties.title = "Micah Willins - Detailed ATS Resume"
    doc.core_properties.author = "Micah Willins"
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Header
    name_p = doc.add_paragraph()
    name_run = name_p.add_run(container.find("h1").get_text(strip=True).upper())
    name_run.font.bold = True
    name_run.font.size = Pt(18)
    name_run.font.color.rgb = BRAND_BLUE
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(clean(container.find("div", class_="title-line").get_text()))
    title_run.font.bold = True
    title_run.font.size = Pt(11)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    contact_items = [clean(s.get_text()) for s in container.select(".top-banner-contact .contact-item")]
    contact_items = [c for c in contact_items if c and c != "|"]
    contact_p = doc.add_paragraph()
    contact_run = contact_p.add_run(" | ".join(contact_items))
    contact_run.font.size = Pt(10)
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links_p = doc.add_paragraph()
    links_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(links_p, "LinkedIn", LINKEDIN_URL)
    separator = links_p.add_run(" | ")
    separator.font.name = "Aptos Narrow"
    separator.font.size = Pt(7.5)
    separator.italic = True
    add_hyperlink(
        links_p,
        "View Hosted Resume for detailed specifics, verified references, and documented award samples",
        RESUME_URL,
    )

    # Sections rendered in source order: Career Summary, Strategic Focus, Technical Skills
    section_headers = container.find_all("div", class_="section-header", recursive=False)
    for header in section_headers:
        title = clean(header.get_text())
        if title in ("Documentation",):
            continue
        if "Professional Experience" in title:
            break
        add_heading(doc, title)
        nxt = header.find_next_sibling()
        if nxt and nxt.name == "dl":
            for index, row in enumerate(nxt.select(".banded-row")):
                para = doc.add_paragraph()
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.keep_together = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEF3FA" if index % 2 == 0 else "FFFFFF")
                para._p.get_or_add_pPr().append(shading)
                label = para.add_run(clean(row.find("dt").get_text()) + ": ")
                label.font.name = "Aptos Display"
                label.bold = True
                label.font.color.rgb = BRAND_BLUE
                para.add_run(clean(row.find("dd").get_text())).font.name = "Aptos Narrow"
        if nxt and nxt.name == "p":
            # Preserve <br> line breaks and bold labels within the paragraph.
            lines: list[list[Tag | NavigableString]] = [[]]
            for child in nxt.children:
                if isinstance(child, Tag) and child.name == "br":
                    lines.append([])
                else:
                    lines[-1].append(child)
            for line_nodes in lines:
                if not line_nodes:
                    continue
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(3)
                for node in line_nodes:
                    if isinstance(node, NavigableString):
                        text = clean(str(node))
                        if text:
                            para.add_run(text).font.name = "Aptos Display"
                    elif isinstance(node, Tag) and node.name in ("b", "strong"):
                        run = para.add_run(clean(node.get_text()))
                        run.font.name = "Aptos Display"
                        run.bold = True

    # Professional Experience
    add_heading(doc, "Professional Experience")
    for details in container.find_all("details"):
        summary = details.find("summary")
        role_p = doc.add_paragraph()
        role_run = role_p.add_run(clean(summary.get_text()))
        role_run.font.name = "Aptos Display"
        role_run.bold = True
        role_run.font.size = Pt(10.5)
        role_p.paragraph_format.space_before = Pt(8)
        role_p.paragraph_format.space_after = Pt(1)
        role_p.paragraph_format.keep_with_next = True

        body = details.find("div")
        lead = body.find(["p"])
        if lead is not None:
            lead_p = doc.add_paragraph()
            lead_run = lead_p.add_run(clean(lead.get_text()))
            lead_run.font.name = "Aptos Narrow"
            lead_run.italic = True
            lead_p.paragraph_format.left_indent = Inches(0.2)
            lead_p.paragraph_format.space_after = Pt(3)
            lead_p.paragraph_format.keep_with_next = True

        outline = body.find("ul", class_="docx-outline")
        if outline is not None:
            for li in outline.find_all("li", recursive=False):
                lead_text, rest_text = li_text(li)
                bullet_p = doc.add_paragraph(style="List Bullet")
                bullet_p.paragraph_format.left_indent = Inches(0.4)
                bullet_p.paragraph_format.space_after = Pt(3)
                bullet_p.paragraph_format.keep_together = True
                if lead_text:
                    lead_run = bullet_p.add_run(lead_text)
                    lead_run.font.name = "Aptos Narrow"
                    lead_run.bold = True
                if rest_text:
                    if lead_text:
                        bullet_p.add_run(" ")
                    bullet_p.add_run(rest_text).font.name = "Aptos Narrow"

    # Education / Awards (verbatim from source)
    for header in section_headers:
        title = clean(header.get_text())
        if title not in ("Education", "Awards"):
            continue
        add_heading(doc, title)
        nxt = header.find_next_sibling()
        if nxt and nxt.name == "p":
            # Rebuild using <br> boundaries directly from the tag.
            text = nxt.decode_contents()
            for chunk in re.split(r"<br\s*/?>", text):
                chunk_text = clean(BeautifulSoup(chunk, "html.parser").get_text())
                if chunk_text:
                    p = doc.add_paragraph(chunk_text, style="List Bullet")
                    for run in p.runs:
                        run.font.name = "Aptos Narrow"
                    p.paragraph_format.space_after = Pt(1)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    build()
