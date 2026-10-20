"""Validate generated resume downloads against the approved HTML source."""

import re
import unittest
from hashlib import sha256
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]


def normalize(text: str) -> str:
    return re.sub(r"\s+", "", text).replace("\u2011", "-")


class ResumeExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = BeautifulSoup(
            (ROOT / "Resume.html").read_text(encoding="utf-8"), "html.parser"
        ).select_one(".container")
        cls.word = Document(ROOT / "Micah-Willins-ATS.docx")
        cls.word_text = normalize("\n".join(p.text for p in cls.word.paragraphs))
        cls.pdf = PdfReader(ROOT / "assets" / "Micah-Willins-Resume.pdf")
        cls.pages = [normalize(page.extract_text()) for page in cls.pdf.pages]

    def test_full_ats_experience_bullets(self):
        bullets = self.source.select("details ul.docx-outline > li")
        self.assertTrue(bullets)
        for bullet in bullets:
            with self.subTest(bullet=bullet.get_text()[:70]):
                self.assertIn(normalize(bullet.get_text()), self.word_text)
        self.assertFalse(self.word.tables)
        self.assertNotIn("COREKEYWORDS", self.word_text)

    def test_focus_and_skill_rows_in_both_downloads(self):
        rows = self.source.select(".banded-row")
        self.assertTrue(rows)
        shaded = [p for p in self.word.paragraphs if p._p.xpath("./w:pPr/w:shd")]
        self.assertEqual(len(shaded), len(rows))
        for row in rows:
            label = normalize(row.dt.get_text())
            value = normalize(row.dd.get_text())
            with self.subTest(category=row.dt.get_text()):
                self.assertIn(label + ":" + value, self.word_text)
                if row.parent.get("id") == "technical-skills":
                    self.assertIn(label, "".join(self.pages))
                    self.assertNotIn(value, "".join(self.pages))
                else:
                    self.assertTrue(any(label in page and value in page for page in self.pages))

    def test_pdf_role_summaries_not_detailed_bullets(self):
        for role in self.source.select("details"):
            title = normalize(role.summary.get_text())
            summary = normalize(role.find("div").find("p").get_text())
            with self.subTest(role=role.summary.get_text()):
                self.assertTrue(any(title in page and summary in page for page in self.pages))
                for bullet in role.select("ul.docx-outline > li"):
                    self.assertNotIn(normalize(bullet.get_text()), "".join(self.pages))
        self.assertEqual(len(self.pdf.pages), 1, "The PDF overview must remain one page")

    def test_intro_education_and_awards(self):
        for header in self.source.select(".section-header"):
            title = header.get_text(strip=True)
            if title not in {"Career Summary", "Education", "Awards"}:
                continue
            for line in re.split(r"<br\s*/?>", header.find_next_sibling("p").decode_contents()):
                text = normalize(BeautifulSoup(line, "html.parser").get_text())
                if text:
                    with self.subTest(section=title, text=text[:60]):
                        self.assertIn(text, self.word_text)
                        if title == "Awards":
                            self.assertNotIn(text, "".join(self.pages))
                        else:
                            self.assertIn(text, "".join(self.pages))
        self.assertNotIn("AWARDS", "".join(self.pages))

    def test_download_copies_identical(self):
        groups = [
            ["Micah-Willins-ATS.docx", "docs/Micah-Willins-ATS.docx",
             "docs/assets/Micah-Willins-ATS.docx"],
            ["assets/Micah-Willins-Resume.pdf", "docs/assets/Micah-Willins-Resume.pdf"],
        ]
        for group in groups:
            self.assertEqual(len({sha256((ROOT / path).read_bytes()).digest() for path in group}), 1)

    def test_online_resume_links(self):
        expected = {
            "https://www.linkedin.com/in/indatawetrust-mw",
            "https://indatawetrust-mw.github.io/MicahWillins/",
        }
        word_links = {
            rel.target_ref for rel in self.word.part.rels.values()
            if rel.reltype.endswith("/hyperlink")
        }
        pdf_links = set()
        for page in self.pdf.pages:
            for reference in page.get("/Annots", []):
                action = reference.get_object().get("/A")
                if action and "/URI" in action:
                    pdf_links.add(str(action["/URI"]))
        self.assertTrue(expected <= word_links)
        self.assertTrue(expected <= pdf_links)
        self.assertIn(
            normalize("View Hosted Resume for detailed specifics, verified references, "
                      "and documented award samples"),
            "".join(self.pages),
        )
        self.assertNotIn(normalize("View Full Details & References"), "".join(self.pages))

    def test_word_hyperlink_text_and_typography(self):
        links = self.word.element.xpath("//w:hyperlink")
        self.assertEqual(
            ["".join(link.xpath(".//w:t/text()")) for link in links],
            ["LinkedIn", "View Hosted Resume for detailed specifics, verified references, "
             "and documented award samples"],
        )
        for link in links:
            for run in link.xpath("./w:r"):
                properties = run.find(qn("w:rPr"))
                self.assertEqual(properties.find(qn("w:rFonts")).get(qn("w:ascii")), "Aptos Narrow")
                self.assertEqual(properties.find(qn("w:sz")).get(qn("w:val")), "15")
                self.assertIsNotNone(properties.find(qn("w:i")))
                self.assertIsNone(properties.find(qn("w:b")))


if __name__ == "__main__":
    unittest.main()
